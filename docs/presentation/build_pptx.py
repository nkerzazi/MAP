# -*- coding: utf-8 -*-
"""Génère la présentation PFE de la plateforme vidéo MAP (charte « Newsroom »)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --- Charte MAP ---
INK      = RGBColor(0x0F, 0x27, 0x42)
INK_600  = RGBColor(0x13, 0x35, 0x5C)
RED      = RGBColor(0xC1, 0x27, 0x2D)
PAPER    = RGBColor(0xF4, 0xF5, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
MUTED    = RGBColor(0x8A, 0x93, 0xA3)
LINE     = RGBColor(0xE3, 0xE6, 0xEA)
TEXT     = RGBColor(0x1F, 0x27, 0x33)
FONT     = "Segoe UI"

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_font(run, size, color, bold=False, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def rect(slide, x, y, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def para(tf, text, size, color, bold=False, italic=False, space_after=6, level=0,
         first=False, bullet=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    if bullet:
        _bullet(p, color)
    run = p.add_run()
    run.text = text
    _set_font(run, size, color, bold, italic, font)
    return p


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    buChar = pPr.makeelement(qn('a:buChar'), {'char': '▪'})
    pPr.append(buFont)
    pPr.append(buChar)


def footer(slide, page):
    rect(slide, 0, SH - Inches(0.34), SW, Inches(0.34), INK)
    tb, tf = textbox(slide, Inches(0.5), SH - Inches(0.34), Inches(9), Inches(0.34),
                     MSO_ANCHOR.MIDDLE)
    para(tf, "MAP — Plateforme de diffusion vidéo interne · PFE 2026", 9, RGBColor(0xA9, 0xBC, 0xD4), first=True)
    tb2, tf2 = textbox(slide, SW - Inches(1.4), SH - Inches(0.34), Inches(0.9), Inches(0.34),
                       MSO_ANCHOR.MIDDLE)
    p = para(tf2, str(page), 9, RGBColor(0xA9, 0xBC, 0xD4), first=True)
    p.alignment = PP_ALIGN.RIGHT


def header(slide, kicker, title):
    rect(slide, 0, 0, SW, Inches(1.15), PAPER)
    rect(slide, 0, 0, Inches(0.18), Inches(1.15), RED)
    tb, tf = textbox(slide, Inches(0.6), Inches(0.18), Inches(12), Inches(0.9))
    para(tf, kicker.upper(), 11, RED, bold=True, space_after=2, first=True)
    para(tf, title, 26, INK, bold=True)
    rect(slide, Inches(0.6), Inches(1.12), Inches(2.0), Pt(3), INK)


def base(slide, fill=WHITE):
    rect(slide, 0, 0, SW, SH, fill)


# ---------------------------------------------------------------- slides
def title_slide():
    s = prs.slides.add_slide(BLANK)
    base(s, INK)
    rect(s, 0, 0, Inches(0.28), SH, RED)
    # logo mark
    rect(s, Inches(0.9), Inches(0.85), Inches(0.34), Inches(0.34), RED)
    tb, tf = textbox(s, Inches(1.4), Inches(0.82), Inches(8), Inches(0.5))
    para(tf, "MAP · MAGHREB ARABE PRESSE", 14, WHITE, bold=True, first=True)

    tb, tf = textbox(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.4))
    para(tf, "Plateforme de gestion et de diffusion", 40, WHITE, bold=True, space_after=2, first=True)
    para(tf, "de contenus vidéo", 40, WHITE, bold=True, space_after=10)
    para(tf, "Solution interne entièrement auto-hébergée — streaming HLS, sans aucun service externe.",
         16, RGBColor(0xA9, 0xBC, 0xD4))

    rect(s, Inches(0.9), Inches(5.4), Inches(4.2), Pt(3), RED)
    tb, tf = textbox(s, Inches(0.9), Inches(5.6), Inches(11), Inches(1.4))
    para(tf, "Projet de Fin d'Études — Youssra Zounaki", 16, WHITE, bold=True, space_after=4, first=True)
    para(tf, "EHEI Oujda · Stage à la MAP, Rabat", 13, RGBColor(0xA9, 0xBC, 0xD4), space_after=2)
    para(tf, "Encadré · 2026", 12, MUTED)


def agenda_slide():
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, "Plan de la présentation", "Sommaire")
    items = [
        ("01", "Problématique", "Contexte MAP, enjeux de souveraineté, objectifs"),
        ("02", "Analyse", "Besoins fonctionnels & non fonctionnels, acteurs"),
        ("03", "Conception", "Architecture, stack, données, pipeline, UI"),
        ("04", "Implémentation", "Backend, frontend, API, tests, déploiement"),
        ("05", "Démonstration & résultats", "Parcours bout-en-bout, captures"),
        ("06", "Perspectives d'amélioration", "Sécurité, vidéo, UX, ops"),
    ]
    y = Inches(1.5)
    for num, t, d in items:
        rect(s, Inches(0.6), y, Inches(0.62), Inches(0.62), INK)
        tbn, tfn = textbox(s, Inches(0.6), y, Inches(0.62), Inches(0.62), MSO_ANCHOR.MIDDLE)
        p = para(tfn, num, 16, WHITE, bold=True, first=True); p.alignment = PP_ALIGN.CENTER
        tb, tf = textbox(s, Inches(1.4), y - Inches(0.02), Inches(11), Inches(0.7), MSO_ANCHOR.MIDDLE)
        para(tf, t, 17, INK, bold=True, space_after=1, first=True)
        para(tf, d, 11.5, MUTED)
        y += Inches(0.92)
    footer(s, len(prs.slides))


def section_slide(num, title, subtitle):
    s = prs.slides.add_slide(BLANK)
    base(s, INK)
    rect(s, 0, Inches(2.6), SW, Inches(2.3), INK_600)
    rect(s, Inches(0.9), Inches(2.6), Inches(0.14), Inches(2.3), RED)
    tb, tf = textbox(s, Inches(1.3), Inches(2.55), Inches(11), Inches(2.4), MSO_ANCHOR.MIDDLE)
    para(tf, num, 18, RED, bold=True, space_after=4, first=True)
    para(tf, title, 38, WHITE, bold=True, space_after=6)
    para(tf, subtitle, 15, RGBColor(0xA9, 0xBC, 0xD4))


def content_slide(kicker, title, blocks, page_note=None):
    """blocks: list of (text, size, color, bold, bullet, level, space_after)."""
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, kicker, title)
    tb, tf = textbox(s, Inches(0.65), Inches(1.5), Inches(12.0), Inches(5.4))
    for i, b in enumerate(blocks):
        text, size, color, bold, bullet, level, sa = b
        para(tf, text, size, color, bold=bold, bullet=bullet, level=level,
             space_after=sa, first=(i == 0))
    footer(s, len(prs.slides))
    return s


def b(text, size=14, color=TEXT, bold=False, bullet=True, level=0, sa=7):
    return (text, size, color, bold, bullet, level, sa)


def two_col_slide(kicker, title, left_title, left_items, right_title, right_items):
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, kicker, title)
    cols = [(Inches(0.65), left_title, left_items), (Inches(7.0), right_title, right_items)]
    for x, ctitle, items in cols:
        card = rect(s, x, Inches(1.55), Inches(5.7), Inches(5.0), PAPER)
        rect(s, x, Inches(1.55), Inches(5.7), Inches(0.55), INK)
        tbh, tfh = textbox(s, x + Inches(0.25), Inches(1.55), Inches(5.2), Inches(0.55), MSO_ANCHOR.MIDDLE)
        para(tfh, ctitle, 15, WHITE, bold=True, first=True)
        tb, tf = textbox(s, x + Inches(0.3), Inches(2.3), Inches(5.1), Inches(4.1))
        for i, it in enumerate(items):
            txt, lvl = it if isinstance(it, tuple) else (it, 0)
            para(tf, txt, 12.5, TEXT, bullet=True, level=lvl, space_after=6, first=(i == 0))
    footer(s, len(prs.slides))


def cards_slide(kicker, title, cards):
    """cards: list of (big, label)."""
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, kicker, title)
    n = len(cards)
    gap = Inches(0.3)
    total_w = SW - Inches(1.3)
    cw = Emu(int((total_w - gap * (n - 1)) / n))
    x = Inches(0.65)
    for big, label in cards:
        rect(s, x, Inches(2.0), cw, Inches(2.4), PAPER, line=LINE)
        rect(s, x, Inches(2.0), cw, Pt(4), RED)
        tb, tf = textbox(s, x, Inches(2.45), cw, Inches(1.8), MSO_ANCHOR.MIDDLE)
        p = para(tf, big, 34, INK, bold=True, first=True); p.alignment = PP_ALIGN.CENTER
        p2 = para(tf, label, 12, MUTED); p2.alignment = PP_ALIGN.CENTER
        x = Emu(x + cw + gap)
    return s


def arch_slide():
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, "Conception", "Architecture générale")
    boxes = [
        ("Angular 17 SPA", "Admin · Éditeur · Visiteur", INK),
        ("API REST /api/v1", "ASP.NET Core 8 · JWT · RBAC", RED),
        ("PostgreSQL 16", "EF Core 8 · full-text", INK),
    ]
    x = Inches(0.65)
    bw = Inches(3.7)
    for t, sub, col in boxes:
        rect(s, x, Inches(1.85), bw, Inches(1.3), col)
        tb, tf = textbox(s, x + Inches(0.15), Inches(1.85), bw - Inches(0.3), Inches(1.3), MSO_ANCHOR.MIDDLE)
        p = para(tf, t, 16, WHITE, bold=True, first=True); p.alignment = PP_ALIGN.CENTER
        p2 = para(tf, sub, 11, RGBColor(0xD7, 0xDE, 0xE8)); p2.alignment = PP_ALIGN.CENTER
        if x < Inches(8):
            ar = textbox(s, x + bw, Inches(1.85), Inches(0.45), Inches(1.3), MSO_ANCHOR.MIDDLE)[1]
            pa = para(ar, "→", 22, MUTED, bold=True, first=True); pa.alignment = PP_ALIGN.CENTER
        x = Emu(x + bw + Inches(0.45))
    # second row
    boxes2 = [
        ("MinIO (S3)", "Originaux + segments HLS"),
        ("Hangfire + FFmpeg 6", "File de transcodage → HLS multi-débit"),
        ("Nginx", "Reverse proxy · diffusion"),
    ]
    x = Inches(0.65)
    for t, sub in boxes2:
        rect(s, x, Inches(3.6), bw, Inches(1.2), PAPER, line=LINE)
        rect(s, x, Inches(3.6), Pt(4), Inches(1.2), RED)
        tb, tf = textbox(s, x + Inches(0.2), Inches(3.6), bw - Inches(0.35), Inches(1.2), MSO_ANCHOR.MIDDLE)
        para(tf, t, 14, INK, bold=True, first=True)
        para(tf, sub, 10.5, MUTED)
        x = Emu(x + bw + Inches(0.45))
    tb, tf = textbox(s, Inches(0.65), Inches(5.1), Inches(12), Inches(1.4))
    para(tf, "Orchestration Docker Compose — pile 100 % auto-hébergée", 14, INK, bold=True, bullet=True, first=True)
    para(tf, "Aucune dépendance externe : pas de YouTube/Vimeo, pas de cloud public, pas de CDN tiers.",
         12.5, TEXT, bullet=True)
    para(tf, "Visionnage anonyme autorisé ; authentification requise pour commenter, liker, partager.",
         12.5, TEXT, bullet=True)
    footer(s, len(prs.slides))


def lifecycle_slide():
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, "Analyse", "Cycle de vie d'une vidéo")
    states = ["Draft", "Processing", "Ready", "Published", "Archived"]
    x = Inches(0.65)
    bw = Inches(2.15)
    for i, st in enumerate(states):
        col = INK if st != "Published" else RED
        rect(s, x, Inches(2.4), bw, Inches(1.0), col)
        tb, tf = textbox(s, x, Inches(2.4), bw, Inches(1.0), MSO_ANCHOR.MIDDLE)
        p = para(tf, st, 15, WHITE, bold=True, first=True); p.alignment = PP_ALIGN.CENTER
        if i < len(states) - 1:
            ar = textbox(s, x + bw, Inches(2.4), Inches(0.35), Inches(1.0), MSO_ANCHOR.MIDDLE)[1]
            pa = para(ar, "→", 20, MUTED, bold=True, first=True); pa.alignment = PP_ALIGN.CENTER
        x = Emu(x + bw + Inches(0.35))
    rect(s, Inches(0.65), Inches(3.9), Inches(3.0), Inches(0.85), PAPER, line=LINE)
    tb, tf = textbox(s, Inches(0.8), Inches(3.9), Inches(2.8), Inches(0.85), MSO_ANCHOR.MIDDLE)
    para(tf, "Failed  (échec pipeline)", 13, RED, bold=True, first=True)
    para(tf, "→ retry automatique (Hangfire)", 11, MUTED)
    tb, tf = textbox(s, Inches(0.65), Inches(5.1), Inches(12), Inches(1.4))
    para(tf, "L'éditeur initialise une vidéo (Draft) et téléverse le fichier en chunks dans MinIO.", 13, TEXT, bullet=True, first=True)
    para(tf, "Un job Hangfire enfile le transcodage ; FFmpeg génère un HLS multi-débit (360p/720p/1080p).", 13, TEXT, bullet=True)
    para(tf, "Succès → Ready (publiable) ; publication → Published (visible au catalogue public).", 13, TEXT, bullet=True)
    footer(s, len(prs.slides))


def image_slide(kicker, title, captions_images):
    s = prs.slides.add_slide(BLANK)
    base(s)
    header(s, kicker, title)
    x = Inches(0.65)
    w = Inches(5.9)
    for cap, img in captions_images:
        if img and os.path.exists(img):
            try:
                s.shapes.add_picture(img, x, Inches(1.7), width=w)
            except Exception:
                pass
        tb, tf = textbox(s, x, Inches(5.4), w, Inches(0.5))
        p = para(tf, cap, 12.5, INK, bold=True, first=True); p.alignment = PP_ALIGN.CENTER
        x = Emu(x + w + Inches(0.3))
    footer(s, len(prs.slides))


# ============================================================== BUILD
title_slide()
agenda_slide()

# --- 01 Problématique
section_slide("01", "Problématique", "Pourquoi une plateforme vidéo interne pour la MAP ?")
content_slide("Problématique · Contexte", "La MAP, agence de presse nationale", [
    b("La Maghreb Arabe Presse produit un volume croissant de contenus vidéo institutionnels.", 14, TEXT),
    b("Ces contenus sont diffusés et archivés sans outil interne dédié, souvent via des plateformes externes.", 14, TEXT),
    b("Enjeu de souveraineté : des contenus officiels et sensibles hébergés hors du contrôle de l'agence.", 14, TEXT),
    b("Besoin d'une chaîne maîtrisée : téléversement → traitement → stockage → diffusion.", 14, INK, bold=True),
])
content_slide("Problématique · Constat", "Les limites de l'existant", [
    b("Dépendance aux services externes (YouTube/Vimeo, cloud public, CDN tiers).", 14, TEXT),
    b("Perte de maîtrise sur la diffusion, les droits et la durée de vie des contenus.", 14, TEXT),
    b("Absence de gestion centralisée des rôles et des accès (qui publie quoi ?).", 14, TEXT),
    b("Pas de statistiques internes ni de traçabilité (audit) des actions.", 14, TEXT),
    b("Contrainte forte du projet : aucune dépendance à un service externe, tout auto-hébergeable.", 14, RED, bold=True),
])
content_slide("Problématique · Objectifs", "Objectifs du projet", [
    b("Plateforme interne, entièrement auto-hébergée, sans aucun service externe.", 14, INK, bold=True),
    b("Téléversement, traitement (transcodage), stockage objet et diffusion en streaming HLS.", 13.5, TEXT),
    b("Contrôle d'accès basé sur les rôles : Admin, Éditeur, Visiteur (RBAC).", 13.5, TEXT),
    b("Recherche du catalogue, statistiques de visionnage et journal d'audit.", 13.5, TEXT),
    b("Interface professionnelle et fluide, bilingue Français / Arabe avec mise en page RTL.", 13.5, TEXT),
    b("Lecteur vidéo développé en interne (hls.js), intégrable et piloté par l'API.", 13.5, TEXT),
])

# --- 02 Analyse
section_slide("02", "Analyse", "Besoins fonctionnels et non fonctionnels")
two_col_slide("Analyse · Rôles", "Besoins fonctionnels par rôle",
    "Admin", [
        "Gestion des utilisateurs et des rôles",
        "Configuration système",
        "Statistiques globales (vues, temps, top)",
        "Journal d'audit",
        "Gestion des catégories",
    ],
    "Éditeur & Visiteur", [
        ("Éditeur", 0),
        ("Téléversement & métadonnées", 1),
        ("Publication / archivage", 1),
        ("Recherche & indexation", 1),
        ("Visiteur", 0),
        ("Navigation du catalogue, visionnage anonyme", 1),
        ("Auth requise : commenter, liker, partager", 1),
    ])
two_col_slide("Analyse · Qualité", "Besoins non fonctionnels",
    "Contraintes clés", [
        "Auto-hébergement total (contrainte forte)",
        "Sécurité : Identity + JWT, RBAC",
        "Async/await systématique (I/O)",
        "Recherche : PostgreSQL full-text (tsvector)",
        "Aucun moteur ni service externe",
    ],
    "Qualité de service", [
        "Performance : HLS multi-débit adaptatif",
        "Scalabilité : jobs asynchrones (Hangfire)",
        "Robustesse : retry du pipeline en cas d'échec",
        "Accessibilité & internationalisation FR/AR + RTL",
        "Erreurs normalisées (ProblemDetails)",
    ])
lifecycle_slide()

# --- 03 Conception
section_slide("03", "Conception", "Architecture, technologies et modèle de données")
arch_slide()
two_col_slide("Conception · Technologies", "Stack technique",
    "Backend & données", [
        "ASP.NET Core 8 (C# 12)",
        "Découpage Domain / Application / Infrastructure / Api",
        "Entity Framework Core 8 + PostgreSQL 16",
        "Hangfire (file de transcodage, storage Postgres)",
        "MinIO (S3) · FFmpeg 6 → HLS",
        "ASP.NET Core Identity + JWT (RBAC)",
    ],
    "Frontend & diffusion", [
        "Angular 17 (standalone, Signals)",
        "Tailwind CSS (design « Newsroom ») + composants headless",
        "Lecteur interne hls.js",
        "i18n FR/AR maison + bascule RTL",
        "Nginx (reverse proxy) · Docker Compose",
        "Tests : xUnit + Testcontainers · Jest + Cypress",
    ])
content_slide("Conception · Données", "Modèle de données (PostgreSQL / EF Core)", [
    b("Identité & RBAC : User, Role, UserRole (mot de passe haché, unicité e-mail).", 13.5, TEXT),
    b("Catalogue : Video (statut, slug, SearchVector), Category, Tag, VideoTag.", 13.5, TEXT),
    b("Diffusion : VideoRendition (360p/720p/1080p, clé manifeste), TranscodeJob (statut, progression, retry).", 13.5, TEXT),
    b("Engagement : Comment, Like (unique par vidéo/utilisateur), Share, ViewEvent (télémétrie).", 13.5, TEXT),
    b("Administration : AuditLog (traçabilité), SystemSetting (configuration clé/valeur).", 13.5, TEXT),
    b("Recherche full-text PostgreSQL (tsvector) sur titre / description / tags — aucun moteur externe.", 13.5, INK, bold=True),
])
content_slide("Conception · Pipeline vidéo", "Du téléversement à la diffusion HLS", [
    b("1 — L'éditeur initialise une vidéo (Draft) et téléverse le fichier en chunks.", 14, TEXT),
    b("2 — Le fichier original est stocké dans MinIO (originals/{videoId}/source).", 14, TEXT),
    b("3 — Un job Hangfire est enfilé ; la vidéo passe en Processing.", 14, TEXT),
    b("4 — FFmpeg génère un HLS multi-débit (360p/720p/1080p) + manifeste maître dans MinIO.", 14, TEXT),
    b("5 — Succès → Ready (publiable) ; échec → Failed + retry Hangfire.", 14, TEXT),
    b("6 — Diffusion : le lecteur consomme le manifeste via un proxy /hls (Nginx + API).", 14, INK, bold=True),
])
two_col_slide("Conception · Interfaces", "Conception des interfaces",
    "Direction « Newsroom »", [
        "Bleu encre + rouge MAP, sobre et éditorial",
        "Inspire la confiance d'une agence officielle",
        "Tailwind + composants accessibles (headless)",
        "État applicatif via Angular Signals",
    ],
    "Trois espaces distincts", [
        "Visiteur : barre haute, catalogue public",
        "Éditeur : back-office (téléversement, suivi)",
        "Admin : back-office (utilisateurs, stats, config)",
        "Bilingue FR/AR — bascule à chaud + RTL global",
    ])

# --- 04 Implémentation
section_slide("04", "Implémentation", "Réalisation, tests et déploiement")
content_slide("Implémentation · Méthode", "Démarche d'ingénierie", [
    b("Flux structuré : brainstorming → spécification → plan d'implémentation → exécution.", 14, TEXT),
    b("Développement piloté par les tests (TDD) : test rouge → code minimal → test vert → commit.", 14, INK, bold=True),
    b("Livraison incrémentale par phases, chacune produisant un logiciel fonctionnel et testé.", 14, TEXT),
    b("Commits fréquents et atomiques ; revue de cohérence vis-à-vis de la spécification.", 14, TEXT),
])
two_col_slide("Implémentation · Backend", "Backend — phases livrées",
    "Phases 1 → 3", [
        "Socle, EF Core, migrations, seed RBAC",
        "Auth : Identity (hash), JWT, gardes par rôle",
        "Catalogue : recherche full-text, publication",
    ],
    "Phases 4 → 6", [
        "Diffusion HLS : /stream + proxy /hls (MinIO)",
        "Engagement : commentaires, likes, partages",
        "Admin & analytique : stats, audit, configuration",
        "+ endpoints Catégories (CRUD) & Tags",
    ])
two_col_slide("Implémentation · Frontend", "Frontend — phases livrées",
    "Socle & espaces", [
        "Phase 1 : socle + design system + espace public",
        "Phase 2 : auth (AuthStore signals) + espace Éditeur",
        "Phase 3 : espace Admin (5 écrans)",
    ],
    "Internationalisation & qualité", [
        "Phase 4 : i18n FR/AR + RTL (service maison)",
        "i18n étendue à tous les écrans éditeur & admin",
        "Upload chunké + suivi de transcodage",
        "Gestion d'erreur (états de chargement robustes)",
    ])
content_slide("Implémentation · API REST", "API REST versionnée (/api/v1)", [
    b("Auth : POST /auth/login · /auth/register · GET /auth/me (JWT Bearer).", 13, TEXT),
    b("Catalogue : GET /videos (recherche, filtres, pagination) · GET /videos/{id} · /stream · /hls/{**}.", 13, TEXT),
    b("Éditeur : POST /videos · /upload/chunk · /upload/complete · PUT /videos/{id} · /publish · /archive.", 13, TEXT),
    b("Engagement : commentaires, likes, partages, /views (télémétrie).", 13, TEXT),
    b("Admin : GET /stats · /audit · GET/PUT /config · /users (+ rôles) · /categories · /tags.", 13, TEXT),
    b("Erreurs normalisées via ProblemDetails ; accès protégés par [Authorize(Roles=…)].", 13, INK, bold=True),
])
cards_slide("Implémentation · Tests & qualité", "Qualité vérifiée par les tests", [
    ("63", "tests backend\n(xUnit + Testcontainers)"),
    ("65", "tests frontend\n(Jest)"),
    ("0", "donnée factice\nen production"),
    ("3", "migrations\nde base de données"),
])
content_slide("Implémentation · Déploiement", "Déploiement auto-hébergé (Docker Compose)", [
    b("Pile complète : PostgreSQL, Redis, MinIO, API, worker (FFmpeg), Nginx.", 14, TEXT),
    b("Migrations appliquées et compte administrateur initialisé au démarrage de l'API.", 14, TEXT),
    b("Tests d'intégration sur conteneurs réels (Postgres + MinIO) — preuve de fonctionnement, pas de simulation.", 14, TEXT),
    b("Aucun service externe : conforme à la contrainte de souveraineté du projet.", 14, INK, bold=True),
])

# --- 05 Démo
section_slide("05", "Démonstration & résultats", "Le produit, en conditions réelles")
fr_img = os.path.join(ROOT, "map-fr.png")
ar_img = os.path.join(ROOT, "map-ar2.png")
image_slide("Démonstration", "Interface — Français (LTR) & Arabe (RTL)", [
    ("Catalogue — Français (gauche → droite)", fr_img),
    ("Catalogue — Arabe (mise en page RTL miroir)", ar_img),
])
content_slide("Démonstration · Bout-en-bout", "Parcours validé sur la pile réelle", [
    b("Connexion de l'administrateur (compte initialisé) → jeton JWT émis et vérifié (/auth/me).", 13.5, TEXT),
    b("Création de catégories via l'API → persistées en PostgreSQL (slug généré, accents normalisés).", 13.5, TEXT),
    b("Inscription d'un éditeur + attribution du rôle Éditeur par l'administrateur.", 13.5, TEXT),
    b("Statistiques calculées en direct depuis la base (nombre d'utilisateurs, de vidéos…).", 13.5, TEXT),
    b("Le frontend consomme l'API réelle ; états de chargement et d'erreur gérés proprement.", 13.5, INK, bold=True),
])

# --- 06 Perspectives
section_slide("06", "Perspectives d'amélioration", "Faire évoluer la plateforme")
two_col_slide("Perspectives · 1/2", "Pistes d'évolution (1/2)",
    "Sécurité & comptes", [
        "Jetons de rafraîchissement (refresh tokens)",
        "Authentification multi-facteurs (MFA)",
        "Politique de mots de passe & verrouillage",
        "Limitation de débit (rate-limiting) anti-abus",
    ],
    "Vidéo & diffusion", [
        "Sous-titres / transcription automatique",
        "Génération automatique de miniatures",
        "Chiffrement HLS (clés) / protection des contenus",
        "Diffusion en direct (live streaming)",
    ])
two_col_slide("Perspectives · 2/2", "Pistes d'évolution (2/2)",
    "Expérience & données", [
        "Recommandations et recherche à facettes",
        "Tableaux de bord analytiques enrichis + export",
        "Notifications (fin de transcodage, publication)",
        "Thème sombre pour le lecteur, accessibilité WCAG",
    ],
    "Exploitation (Ops)", [
        "CI/CD et exécution des tests e2e (Cypress) en pipeline",
        "Supervision (Prometheus / Grafana) & alertes",
        "Haute disponibilité & sauvegardes automatisées",
        "Montée en charge horizontale des workers",
    ])

# --- Conclusion
def conclusion_slide():
    s = prs.slides.add_slide(BLANK)
    base(s, INK)
    rect(s, 0, 0, Inches(0.28), SH, RED)
    tb, tf = textbox(s, Inches(0.9), Inches(1.1), Inches(11.5), Inches(1.2))
    para(tf, "CONCLUSION", 14, RED, bold=True, first=True)
    para(tf, "Une plateforme interne, souveraine et extensible", 30, WHITE, bold=True)
    tb, tf = textbox(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(3.0))
    for t in [
        "Une chaîne complète maîtrisée : téléversement, transcodage HLS, stockage objet et diffusion.",
        "Entièrement auto-hébergée — aucune dépendance à un service externe (souveraineté des contenus).",
        "Sécurisée (JWT + RBAC), internationalisée (FR/AR + RTL) et vérifiée par 128 tests automatisés.",
        "Une base claire et modulaire, prête à accueillir les perspectives d'évolution.",
    ]:
        para(tf, t, 15, RGBColor(0xD7, 0xDE, 0xE8), bullet=True, space_after=12, first=(t.startswith("Une chaîne")))
    rect(s, Inches(0.9), Inches(6.0), Inches(3.5), Pt(3), RED)
    tb, tf = textbox(s, Inches(0.9), Inches(6.2), Inches(11), Inches(0.8))
    para(tf, "Merci de votre attention.", 18, WHITE, bold=True, first=True)

conclusion_slide()

out = os.path.join(HERE, "MAP-PFE-presentation.pptx")
prs.save(out)
print("OK:", out, "·", len(prs.slides), "slides")
